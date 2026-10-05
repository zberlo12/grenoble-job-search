"""Generate Natalie's CL (from template) and copy CV to outputs/."""
import shutil
import copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date
from pathlib import Path

REPO = Path(__file__).parent.parent
OUT  = REPO / "outputs"
OUT.mkdir(exist_ok=True)

MONTHS_FR = {1:"janvier",2:"février",3:"mars",4:"avril",5:"mai",6:"juin",
             7:"juillet",8:"août",9:"septembre",10:"octobre",11:"novembre",12:"décembre"}

CL_PARAS = [
    "C'est un constat clinique simple, mais fondamental, qui m'a orientée vers la santé publique : un patient porteur d'une maladie chronique, bien soigné en phase aiguë à l'hôpital, mais qui sort vers un désert médical, sans accès aux soins médico-sociaux ou aux services sociaux, voit son parcours s'arrêter là où il devrait continuer — et j'ai par ailleurs régulièrement pris en charge des patients pour des états probablement évitables par une prévention adaptée.",
    "Ces réflexions, construites par mes expériences cliniques au CHUGA — gériatrie aiguë, neurovasculaire, pneumologie, urgences adultes —, m'ont orientée vers une spécialisation en prévention et promotion de la santé, avec comme fil conducteur les multiples enjeux de santé publique liés au vieillissement de la population : des enjeux transversaux m'obligeant à penser ensemble ce qui est trop souvent pensé en silo, tout en m'impliquant dans d'autres thématiques et auprès de différentes populations afin d'élargir ma vision, notamment la vaccination HPV en milieu scolaire, les parcours de santé des personnes en situation de migration, la santé environnementale — One Health, arboviroses, vagues de chaleur, entre autres.",
    "Animée par ce regard de pont entre clinique individuelle et santé populationnelle, et par la conviction que c'est à l'ARS que cette approche prend tout son sens, tant dans ses orientations structurantes que dans ses actions territoriales, je souhaite aujourd'hui candidater à ce poste de médecin conseiller.",
    "Ce poste m'attire parce qu'il place le médecin à l'interface entre le sanitaire, le médico-social et le social — là où les décisions deviennent concrètes et ancrées dans la réalité du territoire. La recomposition de l'offre de soins en Isère — restructuration de l'offre hospitalière via le GHT, développement des MSP dans les zones sous-dotées, accompagnement des CPTS, déploiement adapté de la télémédecine — requiert un regard à la fois médical et systémique, capable d'articuler pertinence clinique et réalité territoriale. C'est précisément ce que mon parcours m'a appris à mobiliser.",
    "Contribuer à l'analyse des demandes d'autorisation d'activité de soins dans ce contexte m'intéresse précisément parce que cela demande de déchiffrer la pertinence médicale d'un projet, d'évaluer son ancrage territorial, et de le mettre en dialogue avec les besoins réels d'une population et les ressources disponibles. C'est exactement la posture que mon double parcours clinique et de santé publique m'a appris à tenir — celle d'une professionnelle à l'aise pour appréhender les enjeux médicaux comme les enjeux organisationnels, permettant aux acteurs autour de la table de construire des parcours et des dispositifs structurés et performants.",
    "Cette posture de pont, je l'ai expérimentée concrètement avec IS-ICOPE et PREVAX38. J'ai appris à faire dialoguer les acteurs des trois domaines ainsi que les structures hospitalières et la médecine de ville autour d'un objectif commun, tout en anticipant la pérennisation des dispositifs en mobilisant autant que possible l'offre existante. Par ailleurs, je suis convaincue que placer l'usager au centre est ce qui favorise le mieux ce décloisonnement : c'est lui qui donne aux acteurs une raison concrète et partagée de collaborer. Ma thèse de médecine, centrée sur l'impact perçu des usagers d'IS-ICOPE, me l'a confirmé : la majorité ont exprimé avoir bénéficié pour la première fois d'une vision globale de leur santé — et cette expérience les a rendus plus acteurs de leur propre vieillissement et de leur santé. Ces expériences m'ont permis de tisser des liens avec de nombreux acteurs de l'écosystème isérois issus des trois secteurs, tant institutionnels que de terrain. Le décloisonnement guide désormais mon approche professionnelle, et ce poste me permettra de le mettre en œuvre à une échelle structurante, en m'appuyant sur les outils que l'ARS développe et soutient — CPTS, CLS, DAC, appels à projets, CRT, entre autres.",
    "La participation aux CPOM et aux appels à projets est l'une des dimensions qui m'enthousiasme le plus dans ce poste : contribuer à définir des indicateurs médicaux réellement pertinents et utiles aux équipes de terrain, en lien avec les priorités du SRS et du PAPRAPS 2025-2029, et participer à des évaluations — quantitatives et qualitatives — adaptées aux besoins territoriaux et à la réalité de l'activité.",
    "Les missions d'analyse des EIG et d'inspection sont des rôles qu'on ne développe pas pendant l'internat. Les comités de morbi-mortalité que j'ai appréciés en service clinique m'ont appris à adopter une posture d'analyse rigoureuse, tournée vers la compréhension collective plutôt que la sanction. Ma participation aux campagnes de vaccination en milieu scolaire — où j'ai été confrontée à la gestion d'événements indésirables — m'a permis d'exercer cette même posture dans un contexte de santé publique. La hausse des EIG signalés en région AuRA en 2025 me semble témoigner d'une culture du signalement fondée sur un dialogue fluide et une relation de confiance entre les acteurs de terrain et les institutions. Pour le médecin conseiller, ils sont source d'information pour affiner les indicateurs de suivi des établissements, identifier des besoins non exprimés, et faire émerger des projets ou des actions ancrés dans la réalité du terrain. Ces missions, ainsi que la participation au comité médical des médecins hospitaliers, sont des opportunités de formation que j'accueille volontiers.",
    "Enfin, un stage de six mois au sein de la DD38 à la Veille Sanitaire m'a permis de connaître de l'intérieur la culture de l'institution — une expérience que j'ai beaucoup appréciée, tant sur le plan professionnel que personnel. Je serais ravie de contribuer à cette dynamique.",
    "Je joins à cette candidature mon curriculum vitae ainsi que plusieurs lettres de recommandation écrites par : le Dr Laurence Lorcet (regard opérationnel), le Dr Gaëlle Vareilles (regard de santé publique et de terrain), le Pr Gaëtan Gavazzi (regard clinique et médical), et Mmes Carlyne Berthot et Mélissa Guérin (regard décloisonné, médico-social et social). Le Pr José Labarere, coordinateur local du DES de Santé Publique, est également disponible pour échanger à mon sujet.",
    "Dans l'attente de votre retour, je vous adresse, Monsieur le Directeur et Madame la Directrice Adjointe, mes sincères salutations.",
]


def replace_in_doc(doc, old, new):
    for p in doc.paragraphs:
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)


def generate_cl():
    doc = Document(str(REPO / "templates" / "cl_template.docx"))

    # Header table: name, headline, contacts
    table = doc.tables[0]
    cell0 = table.rows[0].cells[0]
    cell1 = table.rows[0].cells[1]

    for p in cell0.paragraphs:
        for run in p.runs:
            run.text = run.text.replace("ZACHARY BERLO", "NATALIE MOLTA")
            run.text = run.text.replace("{{CL_HEADLINE}}", "Docteure Junior — DES Santé Publique")

    for p in cell1.paragraphs:
        for run in p.runs:
            run.text = run.text.replace("zberlo12@gmail.com", "nmolta@chu-grenoble.fr")
            run.text = run.text.replace("+33 7 66 16 06 01", "07.82.76.45.53")
            if "linkedin.com" in run.text:
                run.text = "Grenoble, France"

    today = date.today()
    date_str = f"Grenoble, le {today.day} {MONTHS_FR[today.month]} {today.year}"

    replace_in_doc(doc, "{{DATE}}", date_str)
    replace_in_doc(doc, "{{COMPANY_ADDRESSEE}}", "ARS Auvergne-Rhône-Alpes — Délégation Départementale de l'Isère")
    replace_in_doc(doc, "{{COMPANY_LOCATION}}", "Grenoble, France")
    replace_in_doc(doc, "{{SUBJECT_LINE}}", "Objet : Candidature au poste de Médecin conseiller — Délégation Départementale de l'Isère")
    replace_in_doc(doc, "Madame, Monsieur,", "Monsieur le Directeur, Madame la Directrice Adjointe,")
    replace_in_doc(doc, "Zachary Berlo", "Natalie Molta")

    # Fill first 5 slots
    slots = ["{{OPENING_PARA}}", "{{BODY_PARA_1}}", "{{BODY_PARA_2}}", "{{BODY_PARA_3}}", "{{CLOSING_PARA}}"]
    for slot, text in zip(slots, CL_PARAS[:5]):
        replace_in_doc(doc, slot, text)

    # Style reference: find any filled body paragraph
    style_ref = None
    for p in doc.paragraphs:
        if CL_PARAS[1][:30] in p.text:
            style_ref = p
            break

    # Insert remaining paragraphs before "Cordialement"
    cordial_elem = None
    for p in doc.paragraphs:
        if "Cordialement" in p.text:
            cordial_elem = p._element
            break

    if cordial_elem is not None:
        parent = cordial_elem.getparent()
        for extra_text in reversed(CL_PARAS[5:]):
            new_p = OxmlElement("w:p")
            new_r = OxmlElement("w:r")
            new_t = OxmlElement("w:t")
            new_t.text = extra_text
            new_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            new_r.append(new_t)
            new_p.append(new_r)
            if style_ref is not None:
                pPr = style_ref._element.find(qn("w:pPr"))
                if pPr is not None:
                    new_p.insert(0, copy.deepcopy(pPr))
            parent.insert(list(parent).index(cordial_elem), new_p)

    out = OUT / "ARS Isère DD38 — Médecin Conseiller_CL.docx"
    doc.save(str(out))
    print(f"CL saved: {out}")


def copy_cv():
    src = REPO / "4. CV nat AR.docx"
    dst = OUT / "ARS Isère DD38 — Médecin Conseiller_CV.docx"
    shutil.copy2(str(src), str(dst))
    print(f"CV copied: {dst}")


if __name__ == "__main__":
    generate_cl()
    copy_cv()
